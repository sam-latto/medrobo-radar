import io
import json
from pathlib import Path

import pytest

import database.db as db
from database import jobs_db
from jobs import boards, collector, functions, matcher
from jobs.resume import ResumeParseError, extract_text

FIXTURES = Path(__file__).parent / "fixtures"
RESUME_TEXT = (
    "Jordan Example\nSenior Product Manager\n\nExperience\nAcme Robotics — Product Manager (2021–2026)\n"
    "• Led roadmap for a surgical navigation platform used in 40 hospitals\n"
    "• Partnered with engineering and clinical teams to launch two FDA-cleared features\n"
    "Beta Consulting — Strategy Consultant (2018–2021)\n• Built market-entry strategy for medtech clients\n"
    "Education\nMBA, Example University\n"
)


def _json(name):
    return json.loads((FIXTURES / name).read_text(encoding="utf-8"))


@pytest.fixture
def jobs_env(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "jobs.db"))
    db.init_db()


# ── Board parsing ─────────────────────────────────────────────────────────────

def test_html_to_text_handles_greenhouse_double_escaping():
    text = boards.html_to_text("&lt;p&gt;Build &amp;amp; ship&lt;/p&gt;&lt;ul&gt;&lt;li&gt;Robots&lt;/li&gt;&lt;/ul&gt;")
    assert text == "Build & ship\n\n• Robots"


def test_parse_greenhouse():
    jobs = boards.parse_greenhouse(_json("greenhouse_jobs.json"), "figureai")
    j = jobs[0]
    assert j["provider"] == "greenhouse" and j["board_token"] == "figureai"
    assert j["company"] == "Figure"
    assert j["url"].startswith("https://job-boards.greenhouse.io/figureai/jobs/")
    assert "<" not in j["description"] and "Figure is an AI robotics company" in j["description"]
    assert j["posted_at"] and len(j["posted_at"]) == 10


def test_parse_lever():
    jobs = boards.parse_lever(_json("lever_postings.json"), "palantir", "Palantir")
    j = jobs[0]
    assert j["url"].startswith("https://jobs.lever.co/palantir/")
    assert j["company"] == "Palantir" and j["title"]
    assert j["location"] and j["employment_type"]
    assert len(j["description"]) > 200 and j["posted_at"]


def test_parse_ashby():
    jobs = boards.parse_ashby(_json("ashby_jobs.json"), "1x", "1X")
    j = jobs[0]
    assert j["url"].startswith("https://jobs.ashbyhq.com/1x/")
    assert j["compensation"] and "$" in j["compensation"]
    assert j["remote"] == 0 and j["description"]


def test_parse_workday_posting():
    url = "https://bostondynamics.wd1.myworkdayjobs.com/en-US/Boston_Dynamics/job/Waltham-MA/Senior-Staff-Product-Manager-Spot_R2410"
    token = "bostondynamics.wd1.myworkdayjobs.com|bostondynamics|Boston_Dynamics"
    p = boards.parse_workday_posting(_json("workday_job.json"), url, token, "Boston Dynamics")
    assert p["title"] == "Senior Staff Product Manager-Spot"
    assert p["provider"] == "workday" and p["company"] == "Boston Dynamics"
    assert "Product Manager for Spot" in p["description"]
    assert boards.parse_workday_posting({}, url, token, None) is None  # closed postings


@pytest.mark.parametrize("url,expected", [
    ("https://job-boards.greenhouse.io/figureai/jobs/4698164006", ("greenhouse", "figureai")),
    ("http://boards.greenhouse.io/diligentrobotics/jobs/7807900003", ("greenhouse", "diligentrobotics")),
    ("https://jobs.lever.co/field-ai/bc148b63", ("lever", "field-ai")),
    ("https://jobs.ashbyhq.com/1x/a5cb058a", ("ashby", "1x")),
    ("https://medtronic.wd1.myworkdayjobs.com/en-US/MedtronicCareers/job/Area-Sales-Manager_R1",
     ("workday", "medtronic.wd1.myworkdayjobs.com|medtronic|MedtronicCareers")),
    ("https://bostondynamics.wd1.myworkdayjobs.com/Boston_Dynamics/job/Waltham-MA/PM_R2410",
     ("workday", "bostondynamics.wd1.myworkdayjobs.com|bostondynamics|Boston_Dynamics")),
    ("https://jobs.smartrecruiters.com/Intuitive/744000111412218-senior-product-manager", None),
])
def test_board_from_url(url, expected):
    assert boards.board_from_url(url) == expected


# ── Board discovery ───────────────────────────────────────────────────────────

def test_slug_candidates_and_name_matching():
    full, all_slugs = collector.slug_candidates("Figure AI")
    assert full == ["figureai", "figure-ai"]
    assert all_slugs == ["figureai", "figure-ai", "figure"]
    assert collector.names_match("Figure", "Figure AI")
    assert collector.names_match("Agility Robotics", "Agility Robotics, Inc.")
    assert not collector.names_match("Figure Eight Federal", "Figure AI")
    assert not collector.names_match("Apptronik", "Figure AI")


def test_discover_boards_verifies_and_caches(jobs_env, monkeypatch):
    calls = []

    def gh_name(token):
        calls.append(("gh", token))
        return {"figure": "Figure", "figureai": None}.get(token)

    monkeypatch.setattr(boards, "greenhouse_board_name", gh_name)
    monkeypatch.setattr(boards, "fetch_lever", lambda token, company=None: (_ for _ in ()).throw(LookupError()))
    monkeypatch.setattr(boards, "fetch_ashby", lambda token, company=None: [] if token != "figure-ai" else [{"x": 1}])

    found = collector.discover_boards("Figure AI")
    assert found == [("greenhouse", "figure"), ("ashby", "figure-ai")]
    assert {(b["provider"], b["token"], b["company"]) for b in jobs_db.list_boards()} == {
        ("greenhouse", "figure", "Figure AI"), ("ashby", "figure-ai", "Figure AI")}
    calls.clear()
    assert collector.discover_boards("Figure AI") == []  # checked recently: no HTTP calls
    assert calls == []


# ── Functions and filtering ───────────────────────────────────────────────────

@pytest.mark.parametrize("title,expected", [
    ("Senior Product Manager, Robotics", ["Product Management"]),
    ("Product Marketing Manager", []),
    ("Strategy & Corporate Development Manager", ["Strategy"]),
    ("Chief of Staff to the CEO", ["Strategy"]),
    ("Equipment Maintenance Technician", []),
])
def test_matching_functions(title, expected):
    assert functions.matching_functions(title, ["Product Management", "Strategy"]) == expected


def test_custom_function_keywords_cached(jobs_env, monkeypatch):
    calls = []

    def fake_call_tool(*a, **k):
        calls.append(1)
        return {"keywords": ["Venture Partner", "investment associate"]}

    monkeypatch.setattr(functions, "call_tool", fake_call_tool)
    assert "venture partner" in functions.keywords_for("Venture Capital")
    assert "venture capital" in functions.keywords_for("Venture Capital")
    assert len(calls) == 1
    assert functions.matching_functions("Investment Associate, Robotics Fund", ["Venture Capital"]) == ["Venture Capital"]


@pytest.mark.parametrize("job,locations,include_remote,remote_only,ok", [
    ({"location": "Waltham, MA"}, ["Boston", "MA"], True, False, True),
    ({"location": "San Carlos, CA"}, ["MA"], True, False, False),
    ({"location": "Toronto, Canada"}, ["CA"], True, False, False),  # 'CA' is a whole word, not 'Canada'
    ({"location": "Remote - US", "remote": None}, ["Boston"], True, False, True),
    ({"location": "Remote - US"}, ["Boston"], False, False, False),
    ({"location": "Waltham, MA", "remote": 0}, [], True, True, False),
    ({"location": "Anywhere", "remote": 1}, [], True, True, True),
    ({"location": None, "provider": "web"}, ["Boston"], True, False, True),
])
def test_location_ok(job, locations, include_remote, remote_only, ok):
    assert matcher.location_ok(job, locations, include_remote, remote_only) is ok


def _posting(i, title, company="Acme", location="Boston, MA", industry="healthcare_robotics", **kw):
    return {"url": f"https://ex.com/{i}", "provider": "greenhouse", "board_token": "acme", "company": company,
            "title": title, "location": location, "description": f"Description {i}", "posted_at": f"2026-09-{i:02d}", **kw}


def test_candidate_jobs_filters_and_dedups(jobs_env):
    jobs_db.upsert_posting(_posting(2, "Senior Product Manager", provider="web"), "healthcare_robotics")
    jobs_db.upsert_posting(_posting(1, "Senior Product Manager"), "healthcare_robotics")  # same job, from the board
    jobs_db.upsert_posting(_posting(3, "Mechanical Engineer"), "healthcare_robotics")
    jobs_db.upsert_posting(_posting(4, "Strategy Lead", location="Austin, TX"), "healthcare_robotics")
    jobs_db.upsert_posting(_posting(5, "Product Manager", industry="physical_ai"), "physical_ai")
    out = matcher.candidate_jobs(["healthcare_robotics"], ["Product Management", "Strategy"], ["Boston"])
    assert [(j["title"], j["provider"]) for j in out] == [("Senior Product Manager", "greenhouse")]
    assert out[0]["functions"] == ["Product Management"]
    both = matcher.candidate_jobs(["healthcare_robotics", "physical_ai"], ["Product Management"], [])
    assert {j["url"] for j in both} == {"https://ex.com/1", "https://ex.com/5"}  # board copy kept over web dup


def test_deactivate_missing(jobs_env):
    for i in (1, 2):
        jobs_db.upsert_posting(_posting(i, "Product Manager"), "healthcare_robotics")
    assert jobs_db.deactivate_missing("greenhouse", "acme", {"https://ex.com/1"}) == 1
    assert [p["url"] for p in jobs_db.list_active_postings()] == ["https://ex.com/1"]
    jobs_db.upsert_posting(_posting(2, "Product Manager"), "physical_ai")  # reappears -> active again
    active = {p["url"]: p for p in jobs_db.list_active_postings()}
    assert active["https://ex.com/2"]["industry_ids"] == ["healthcare_robotics", "physical_ai"]


# ── Scoring and advice ────────────────────────────────────────────────────────

def test_score_jobs_batches_caches_and_fixes_company(jobs_env, monkeypatch):
    ids = [jobs_db.upsert_posting(_posting(i, f"Product Manager {i}", company=None if i == 1 else "Acme"),
                                  "healthcare_robotics") for i in range(1, 11)]
    resume = {"id": jobs_db.save_resume("r.txt", RESUME_TEXT), "text": RESUME_TEXT}
    calls = []

    def fake_call_tool(prompt, tool, max_tokens=4096, system=None):
        calls.append({"prompt": prompt, "system": system})
        job_ids = [int(x.split("]")[0]) for x in prompt.split("[job_id=")[1:]]
        return {"scores": [{"job_id": j, "score": 150 if j == ids[0] else 70, "summary": "ok",
                            "strengths": ["a", "b", "c", "d"], "gaps": [], "level_fit": "match",
                            "company": "Figure", "industry_match": j != ids[2]} for j in job_ids]}

    monkeypatch.setattr(matcher, "call_tool", fake_call_tool)
    jobs = matcher.candidate_jobs(["healthcare_robotics"], ["Product Management"], [])
    assert matcher.score_jobs(resume, jobs, ["Healthcare Robotics"]) == 10
    assert len(calls) == 2  # batches of 8
    assert "targeting roles in: Healthcare Robotics" in calls[0]["prompt"]
    assert calls[0]["system"][0]["cache_control"] == {"type": "ephemeral"}
    assert "<resume>" in calls[0]["system"][0]["text"] and "Jordan Example" in calls[0]["system"][0]["text"]
    matches = jobs_db.get_matches(resume["id"])
    assert matches[ids[0]]["score"] == 100  # clamped
    assert len(matches[ids[1]]["strengths"]) == 3
    assert (matches[ids[1]]["industry_match"], matches[ids[2]]["industry_match"]) == (1, 0)
    assert jobs_db.get_posting(ids[0])["company"] == "Figure"   # filled in from the posting
    assert jobs_db.get_posting(ids[1])["company"] == "Acme"     # known company untouched
    assert matcher.score_jobs(resume, jobs, ["Healthcare Robotics"]) == 0 and len(calls) == 2  # already scored


def test_score_jobs_prefers_board_postings_and_ignores_placeholder_company(jobs_env, monkeypatch):
    web_id = jobs_db.upsert_posting(_posting(1, "Product Manager A", company=None, provider="web"), "healthcare_robotics")
    board_id = jobs_db.upsert_posting(_posting(2, "Product Manager B", company=None), "healthcare_robotics")
    resume = {"id": jobs_db.save_resume(None, RESUME_TEXT), "text": RESUME_TEXT}
    monkeypatch.setattr(matcher, "MAX_SCORED_PER_SEARCH", 1)
    seen = []

    def fake_call_tool(prompt, tool, max_tokens=4096, system=None):
        job_ids = [int(x.split("]")[0]) for x in prompt.split("[job_id=")[1:]]
        seen.extend(job_ids)
        return {"scores": [{"job_id": j, "score": 60, "summary": "s", "strengths": [], "gaps": [],
                            "level_fit": "match", "industry_match": True, "company": "Unknown"} for j in job_ids]}

    monkeypatch.setattr(matcher, "call_tool", fake_call_tool)
    matcher.score_jobs(resume, matcher.candidate_jobs(["healthcare_robotics"], ["Product Management"], []), ["X"])
    assert seen == [board_id]  # the single scoring slot goes to the full board posting
    assert jobs_db.get_posting(board_id)["company"] is None  # "Unknown" isn't a company name
    assert web_id not in jobs_db.get_matches(resume["id"])


def test_tailor_resume_flags_unverified_bullets_and_caches(jobs_env, monkeypatch):
    job_id = jobs_db.upsert_posting(_posting(1, "Senior Product Manager"), "healthcare_robotics")
    resume = {"id": jobs_db.save_resume(None, RESUME_TEXT), "text": RESUME_TEXT}
    calls = []

    def fake_call_tool(prompt, tool, max_tokens=4096, system=None):
        calls.append(prompt)
        assert "Never invent" in system[0]["text"]
        return {"fit_overview": "Good", "summary_rewrite": "PM", "keywords_to_add": [], "emphasize": [],
                "de_emphasize": [], "gaps": [], "cover_letter_angle": "x",
                "bullet_rewrites": [
                    {"original": "Led roadmap for a surgical navigation platform used in 40 hospitals",
                     "suggested": "Owned roadmap…", "why": "impact"},
                    {"original": "Managed a $10M P&L", "suggested": "…", "why": "made up"},
                ]}

    monkeypatch.setattr(matcher, "call_tool", fake_call_tool)
    out = matcher.tailor_resume(resume, jobs_db.get_posting(job_id))
    assert [r["verified"] for r in out["bullet_rewrites"]] == [True, False]
    assert matcher.tailor_resume(resume, jobs_db.get_posting(job_id)) == out
    assert len(calls) == 1


def test_overall_recommendations_cached_per_functions(jobs_env, monkeypatch):
    resume = {"id": jobs_db.save_resume(None, RESUME_TEXT), "text": RESUME_TEXT}
    calls = []
    monkeypatch.setattr(matcher, "call_tool", lambda *a, **k: calls.append(1) or {"positioning": "p"})
    matcher.overall_recommendations(resume, ["Strategy", "Product Management"], ["X"], ["a"], [])
    matcher.overall_recommendations(resume, ["Product Management", "Strategy"], ["X"], ["a"], [])
    matcher.overall_recommendations(resume, ["Product Management"], ["X"], ["a"], [])
    assert len(calls) == 2


# ── Resume parsing ────────────────────────────────────────────────────────────

def test_extract_text_txt_and_docx():
    assert "Jordan Example" in extract_text("resume.txt", RESUME_TEXT.encode())
    from docx import Document
    doc = Document()
    for line in RESUME_TEXT.splitlines():
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Skills", "Roadmapping, SQL"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("resume.docx", buf.getvalue())
    assert "surgical navigation platform" in text and "Skills | Roadmapping, SQL" in text


def test_extract_text_rejects_bad_input():
    with pytest.raises(ResumeParseError):
        extract_text("resume.png", b"...")
    with pytest.raises(ResumeParseError):
        extract_text("resume.txt", b"too short")


@pytest.mark.parametrize("provider,url,expected", [
    ("greenhouse", "http://job-boards.greenhouse.io/diligentrobotics/jobs/7807900003", "7807900003"),
    ("greenhouse", "https://boards.greenhouse.io/acme/jobs/123?gh_src=x", "123"),
    ("greenhouse", "https://acme.com/careers?gh_jid=456", "456"),
    ("lever", "https://jobs.lever.co/field-ai/bc148b63-61e4", "bc148b63-61e4"),
    ("ashby", "https://jobs.ashbyhq.com/1x/a5cb058a", "a5cb058a"),
])
def test_posting_id_from_url(provider, url, expected):
    assert collector.posting_id_from_url(provider, url) == expected


def test_search_web_stores_only_found_postings(jobs_env, monkeypatch):
    import requests
    from profiles.loader import load_profile

    results = [
        {"url": "https://job-boards.greenhouse.io/acme/jobs/2", "title": "PM", "content": "c"},
        {"url": "https://jj.wd5.myworkdayjobs.com/en-US/JJ/job/Technical-Product-Owner_R-1",
         "title": "Technical Product Owner - Myworkdayjobs.com", "raw_content": "Full JJ text"},
        {"url": "https://bank.wd1.myworkdayjobs.com/Careers/job/Strategy-Manager_R9",
         "title": "Strategy Manager", "content": "bank"},
        {"url": "https://jobs.smartrecruiters.com/Intuitive/1-senior-product-manager",
         "title": "Senior Product Manager", "raw_content": "Intuitive PM role"},
    ]

    class FakeTavily:
        def __init__(self, api_key): pass
        def search(self, **kw): return {"results": results}

    board_jobs = [{"url": f"https://job-boards.greenhouse.io/acme/jobs/{i}", "provider": "greenhouse",
                   "board_token": "acme", "company": "Acme", "title": f"Job {i}", "description": "d"} for i in (1, 2, 3)]

    def fake_workday(url, company=None):
        if "jj." in url:
            resp = requests.Response(); resp.status_code = 403
            raise requests.HTTPError("403 Forbidden", response=resp)
        return {"url": url, "provider": "workday", "board_token": None, "company": company,
                "title": "Strategy Manager", "description": "bank strategy"}

    monkeypatch.setattr(collector, "TavilyClient", FakeTavily)
    monkeypatch.setattr(collector.boards, "fetch_board", lambda provider, token, company=None: [dict(j) for j in board_jobs])
    monkeypatch.setattr(collector.boards, "fetch_workday_posting", fake_workday)

    profile = load_profile("healthcare_robotics")
    stored = collector.search_web(profile, ["Product Management"], {"Johnson & Johnson"}, lambda m: None)
    urls = {p["url"]: p for p in jobs_db.list_active_postings(["healthcare_robotics"])}
    assert stored == 4
    assert "https://job-boards.greenhouse.io/acme/jobs/2" in urls
    assert "https://job-boards.greenhouse.io/acme/jobs/1" not in urls  # rest of the board not imported
    jj = urls["https://jj.wd5.myworkdayjobs.com/en-US/JJ/job/Technical-Product-Owner_R-1"]
    assert jj["provider"] == "web" and jj["description"] == "Full JJ text"  # 403 -> search text kept
    assert urls["https://jobs.smartrecruiters.com/Intuitive/1-senior-product-manager"]["title"] == "Senior Product Manager"
    assert jobs_db.list_boards() == []  # no stray boards registered for unrelated companies
