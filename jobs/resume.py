"""Turn an uploaded resume (PDF, Word, or text) into plain text."""
import io
import re
from pathlib import Path

MAX_RESUME_CHARS = 30_000


class ResumeParseError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # resumes often lay out sections in tables
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif suffix in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ResumeParseError(f"Unsupported file type '{suffix}'. Upload a PDF, .docx, or .txt file.")
    return clean_text(text)


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text).strip()
    if len(text) < 200:
        raise ResumeParseError(
            "Couldn't read enough text from that resume (is it a scanned image?). "
            "Try another format or paste the text instead."
        )
    return text[:MAX_RESUME_CHARS]
